# -*- coding: utf-8 -*-
"""Shared layout and styling for the CV builders.

Both build_cv.py (software/engineering) and build_cv_hospitality.py render
through CvBuilder so the two CVs stay visually identical — only their content
differs. Nothing in here is CV-specific.

Deps: pip install python-docx qrcode pillow
"""
from pathlib import Path

import qrcode
from qrcode.constants import ERROR_CORRECT_M

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
PORTFOLIO_URL = "https://jamesdillon.uk"
QR_PATH = HERE / "portfolio-qr.png"

NAVY = RGBColor(0x1B, 0x36, 0x5D)
GREY = RGBColor(0x55, 0x5B, 0x66)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
FAINT = RGBColor(0xA0, 0xA8, 0xB4)
NAVY_HEX = "1B365D"
DARK_HEX = "1A1A1A"
RULE = "9BA6B8"
FONT = "Calibri"
SEP = "   •   "


def build_qr(url=PORTFOLIO_URL, path=QR_PATH):
    """Regenerate the portfolio QR. Box size is deliberately large so the code
    stays crisp when the DOCX is printed."""
    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_M, box_size=24, border=2)
    qr.add_data(url)
    qr.make(fit=True)
    qr.make_image(fill_color="black", back_color="white").convert("RGB").save(path)
    return path


class CvBuilder:
    """A single-column A4 CV: ruled section headings, right-aligned dates,
    bulleted entries, and a QR block in the header."""

    def __init__(self):
        self.doc = Document()

        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)   # A4
        sec.left_margin = sec.right_margin = Inches(0.45)
        sec.top_margin = Inches(0.4)
        sec.bottom_margin = Inches(0.38)
        self.content_w = sec.page_width - sec.left_margin - sec.right_margin

        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        normal.font.color.rgb = DARK
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = normal.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0

        # Vertical rhythm. Tighten these to pull a CV onto fewer pages.
        self.sp_heading_before = 11
        self.sp_heading_after = 4.5
        self.sp_role_before = 7
        self.sp_bullet_after = 3.4
        self.sp_skill_after = 3.6

    # ---------- primitives ----------

    @staticmethod
    def set_spacing(p, before=0, after=0, line=1.0):
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = line
        return p

    @staticmethod
    def run(p, text, size=10, bold=False, italic=False, color=DARK):
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
        return r

    @staticmethod
    def hyperlink(p, text, url, size=10, bold=False, color=NAVY_HEX):
        """Insert real external hyperlink. Without this Word autoformats bare
        domains like "jamesdillon.uk" into a relative path, which resolves to a
        local file:/// URL instead of the site."""
        r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)

        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        # child order below is fixed by the CT_RPr schema — do not reorder
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        rPr.append(fonts)
        if bold:
            rPr.append(OxmlElement("w:b"))
        col = OxmlElement("w:color")
        col.set(qn("w:val"), color)
        rPr.append(col)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rPr.append(sz)
        r.append(rPr)

        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)

        link.append(r)
        p._p.append(link)
        return link

    @staticmethod
    def bottom_border(p, color=RULE, size=6, space=2):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:space"), str(space))
        b.set(qn("w:color"), color)
        bdr.append(b)
        pPr.append(bdr)

    @staticmethod
    def keep_with_next(p):
        pPr = p._p.get_or_add_pPr()
        pPr.append(OxmlElement("w:keepNext"))
        return p

    @staticmethod
    def no_cell_margins(cell):
        tcPr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for side in ("left", "right", "top", "bottom"):
            e = OxmlElement("w:" + side)
            e.set(qn("w:w"), "0")
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tcPr.append(mar)

    # ---------- blocks ----------

    def heading(self, text, first=False):
        p = self.doc.add_paragraph()
        self.set_spacing(p, before=0 if first else self.sp_heading_before,
                         after=self.sp_heading_after)
        r = self.run(p, text.upper(), size=11, bold=True, color=NAVY)
        r.font.spacing = Pt(0.9)
        self.bottom_border(p)
        self.keep_with_next(p)
        return p

    def role(self, title, meta, date):
        """Entry title on the left, date right-aligned on the same line."""
        p = self.doc.add_paragraph()
        self.set_spacing(p, before=self.sp_role_before, after=2.5)
        p.paragraph_format.tab_stops.add_tab_stop(self.content_w, WD_TAB_ALIGNMENT.RIGHT)
        self.run(p, title, size=10.5, bold=True, color=NAVY)
        if meta:
            self.run(p, "  |  ", size=10, color=FAINT)
            self.run(p, meta, size=10, color=DARK)
        p.add_run("\t")
        self.run(p, date, size=9, italic=True, color=GREY)
        self.keep_with_next(p)
        return p

    def bullet(self, parts, indent=0.0):
        """parts: str, or list of (text, bold) tuples."""
        p = self.doc.add_paragraph()
        self.set_spacing(p, before=0, after=self.sp_bullet_after, line=1.02)
        p.paragraph_format.left_indent = Inches(indent + 0.17)
        p.paragraph_format.first_line_indent = Inches(-0.17)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(indent + 0.17))
        r = p.add_run("•\t")
        r.font.name = FONT
        r.font.size = Pt(10)
        r.font.color.rgb = NAVY
        r.bold = True
        if isinstance(parts, str):
            parts = [(parts, False)]
        for text, bold in parts:
            self.run(p, text, size=10, bold=bold)
        return p

    def skill_line(self, label, value, label_w=1.3):
        p = self.doc.add_paragraph()
        self.set_spacing(p, before=0, after=self.sp_skill_after, line=1.02)
        p.paragraph_format.left_indent = Inches(label_w)
        p.paragraph_format.first_line_indent = Inches(-label_w)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(label_w))
        self.run(p, label, size=10, bold=True, color=NAVY)
        p.add_run("\t")
        self.run(p, value, size=10)
        return p

    def paragraph(self, text, before=2, after=0, line=1.08):
        p = self.set_spacing(self.doc.add_paragraph(), before=before, after=after, line=line)
        self.run(p, text)
        return p

    def header(self, name, tagline, contact_line, link_line, qr_caption,
               qr_path=QR_PATH, qr_url=PORTFOLIO_URL, qr_label="jamesdillon.uk"):
        """Name and contact details on the left, QR block flush right.

        contact_line / link_line are lists of (text, url_or_None) pairs, joined
        by bullet separators. qr_caption is a list of short lines shown beside
        the QR code.
        """
        qr_w = Inches(1.12)
        right_w = Inches(2.72)
        name_w = self.content_w - right_w

        head = self.doc.add_table(rows=1, cols=2)
        head.alignment = WD_TABLE_ALIGNMENT.LEFT
        head.autofit = False
        head.columns[0].width = name_w
        head.columns[1].width = right_w
        left, right = head.rows[0].cells
        left.width = name_w
        right.width = right_w
        self.no_cell_margins(left)
        self.no_cell_margins(right)

        p = left.paragraphs[0]
        self.set_spacing(p, after=1)
        r = self.run(p, name, size=25, bold=True, color=NAVY)
        r.font.spacing = Pt(0.5)

        p = self.set_spacing(left.add_paragraph(), before=0, after=5)
        self.run(p, tagline, size=10, color=GREY)

        for idx, line in enumerate([contact_line, link_line]):
            if not line:
                continue
            p = self.set_spacing(left.add_paragraph(), before=0, after=1.5 if idx == 0 else 0)
            for i, (text, url) in enumerate(line):
                if i:
                    self.run(p, SEP, size=9, color=FAINT)
                if url:
                    self.hyperlink(p, text, url, size=9, color=DARK_HEX)
                else:
                    self.run(p, text, size=9)

        # QR block: label + caption sit to the LEFT of the QR, which stays flush right
        qr_txt_w = right_w - qr_w - Inches(0.1)
        qr_inner = right.add_table(rows=1, cols=2)
        qr_inner.autofit = False
        qr_txt_cell, qr_img_cell = qr_inner.rows[0].cells
        qr_inner.columns[0].width = qr_txt_w
        qr_inner.columns[1].width = qr_w + Inches(0.1)
        qr_txt_cell.width = qr_txt_w
        qr_img_cell.width = qr_w + Inches(0.1)
        self.no_cell_margins(qr_txt_cell)
        self.no_cell_margins(qr_img_cell)

        p = qr_txt_cell.paragraphs[0]
        self.set_spacing(p, before=3, after=2.5, line=0.95)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self.hyperlink(p, qr_label, qr_url, size=13.5, bold=True)

        for line in qr_caption:
            p = self.set_spacing(qr_txt_cell.add_paragraph(), before=0, after=0, line=0.95)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            self.run(p, line, size=8, color=GREY)

        p = qr_img_cell.paragraphs[0]
        self.set_spacing(p, after=0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture(str(qr_path), width=qr_w)

        # drop the empty placeholder paragraph the outer cell started with
        right.paragraphs[0]._p.getparent().remove(right.paragraphs[0]._p)

        self.set_spacing(self.doc.add_paragraph(), before=0, after=3)

    def footer_note(self, lead, label=None, url=PORTFOLIO_URL, before=9):
        p = self.set_spacing(self.doc.add_paragraph(), before=before, after=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.run(p, lead, size=8.5, color=GREY)
        if label:
            self.hyperlink(p, label, url, size=8.5, bold=True)
        return p

    def save(self, path):
        self.doc.save(path)
        print("saved", path)
